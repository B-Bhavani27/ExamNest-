import csv
import io

from django.http import HttpResponse

import openpyxl
from openpyxl.styles import Font, PatternFill

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm


def export_report(headers, rows, filename, export_format, title=""):
    """
    headers: list of column header strings
    rows: list of lists/tuples — each inner list is one row, same order as headers
    filename: without extension, e.g. "student_report"
    export_format: 'csv' | 'xlsx' | 'pdf'
    title: optional heading printed at the top of the PDF
    """

    export_format = (export_format or "csv").lower()

    if export_format == "xlsx":
        return _export_xlsx(headers, rows, filename, title)

    if export_format == "pdf":
        return _export_pdf(headers, rows, filename, title)

    return _export_csv(headers, rows, filename)


def _export_csv(headers, rows, filename):

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'

    writer = csv.writer(response)
    writer.writerow(headers)

    for row in rows:
        writer.writerow(row)

    return response


def _export_xlsx(headers, rows, filename, title):

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="7C3AED", end_color="7C3AED", fill_type="solid")

    sheet.append(headers)

    for col_index in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_index)
        cell.font = header_font
        cell.fill = header_fill

    for row in rows:
        sheet.append(list(row))

    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max_length + 4, 45)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'

    return response


def _export_pdf(headers, rows, filename, title):

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()
    elements = []

    if title:
        elements.append(Paragraph(title, styles["Title"]))
        elements.append(Spacer(1, 14))

    table_data = [headers] + [list(row) for row in rows]

    table = Table(table_data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#7C3AED")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E5E5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FB")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'

    return response
